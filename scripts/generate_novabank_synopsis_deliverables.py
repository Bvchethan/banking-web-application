from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak


OUTPUT_DIR = Path("output/synopsis")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUTPUT_DIR / "NovaBank_Project_Synopsis_Final.docx"
PDF_PATH = OUTPUT_DIR / "NovaBank_Project_Synopsis_Final.pdf"


TITLE = "NovaBank: A Secure Full-Stack Digital Banking Web Application with AI-Inspired Fraud Detection and Financial Health Analytics"

DETAILS = [
    ("Name", "Chethan BV"),
    ("USN", "241VMTR01537"),
    ("Elective", "Full Stack Development"),
    ("Date of Submission", "25/05/2026"),
]

SECTIONS = [
    ("Problem Statement:", [
        "The rapid digitization of financial services has significantly transformed customer expectations regarding banking accessibility, transparency, and security. However, many digital banking platforms, particularly those designed for small-scale and mid-scale deployments, continue to face challenges such as fragmented user experiences, insufficient role-based access control, weak transaction monitoring, and limited visibility into suspicious fund transfer activities. Customers often lack a unified platform to manage beneficiaries, review transaction histories, perform secure fund transfers, and access meaningful financial insights within a single responsive application.",
        "On the administrative side, many banking systems fail to provide centralized dashboards for monitoring customer activity, identifying high-risk transactions, and maintaining effective operational oversight. Security mechanisms such as encrypted password storage, JWT-based stateless authentication, and properly enforced role-based API protection are often either absent or inadequately implemented.",
        "This project proposes the development of a production-style full-stack digital banking web application titled NovaBank, designed to address these limitations. The system will be built using React and Tailwind CSS for the frontend, and Spring Boot, Spring Security, JWT, Spring Data JPA, and MySQL for the backend. The proposed application will support customer and admin roles, providing secure authentication, account management, beneficiary management, fund transfers with atomic transaction recording, AI-inspired fraud risk scoring, financial health insights, and a centralized administrative dashboard. The study will evaluate the effectiveness of layered security architecture, transactional consistency, and intelligent analytics in a modern web-based banking platform.",
    ]),
    ("Objectives of the Project:", [
        "1. To develop a role-based digital banking platform that enables customers to manage their profiles, beneficiaries, account balances, and fund transfers, while allowing administrators to monitor customer records, transactions, fraud signals, and platform-wide operational metrics through a centralized dashboard.",
        "2. To implement secure authentication and authorization using JWT, BCrypt password encryption, and Spring Security so that only verified and authorized users can access protected banking operations according to their assigned roles.",
        "3. To ensure reliable and consistent fund transfer processing by updating sender and receiver account balances atomically within a single transaction, while maintaining complete dual-entry transaction records for both parties for auditability and transparency.",
        "4. To incorporate an AI-inspired fraud detection module that evaluates each transfer against risk criteria such as high transaction amounts, rapid successive transfers, and newly added beneficiaries, classifying transactions as LOW, MEDIUM, or HIGH risk.",
        "5. To develop a financial health analytics module that generates monthly savings insights, spending trends, and smart suggestions to improve customer financial awareness.",
    ]),
    ("Project Methodology:", [
        "Type of Project: This is an application-based project. NovaBank will be developed as a deployable, production-style full-stack web application rather than a theoretical or purely research-oriented study. The project will cover major software development lifecycle activities including requirement analysis, database design, backend API development, frontend integration, security implementation, testing, and deployment preparation.",
        "Technology Stack: The frontend will be developed using React 18, Tailwind CSS, Axios, and React Router. The backend will be developed using Java Spring Boot 3, Spring Security, JWT, Spring Data JPA, and Maven. MySQL will be used as the relational database. Swagger UI through SpringDoc OpenAPI will be used for API documentation and testing.",
        "System Architecture: The project will follow a layered client-server architecture. The React frontend will function as a Single Page Application and communicate with the Spring Boot backend through secured REST APIs. The backend will be organized into distinct packages such as controller, service, repository, dto, entity, config, security, and exception, ensuring separation of concerns, maintainability, and enterprise-style code organization. A JWT authentication filter will be used to protect all non-public API routes by validating tokens before business logic is executed.",
        "Data Collection Methods: Secondary data will be used to understand digital banking workflows, security practices, and transaction monitoring requirements from existing banking platforms and technical references. Primary data will be simulated through user interactions within the application, including registration, login, beneficiary creation, fund transfers, and transaction review, which will guide the API design and relational database schema.",
    ]),
    ("Key Modules to be Developed:", [
        "1. User Management: Customer registration and admin access/login, JWT-based stateless authentication, BCrypt password encryption, and role-based access control will be implemented to secure the application. Newly registered customers will automatically be provisioned with a banking account upon successful registration.",
        "2. Profile and Account Management: Customers will be able to view and update their personal details such as full name, phone number, and address. The account summary module will provide real-time access to account number, available balance, and currency.",
        "3. Beneficiary Management: Customers will be able to add, edit, and delete trusted beneficiaries. Each beneficiary will store recipient name, account number, bank name, IFSC code, and a personal nickname for convenient use during transfers.",
        "4. Secure Fund Transfer: The system will support account-to-account fund transfers with atomic balance updates and proper transactional consistency. Each transfer will create separate debit and credit transaction entries for sender and receiver, maintaining a complete and auditable transaction history for both parties.",
        "5. Transaction History and Search: Customers will be able to search and review their transaction history by description or reference number. Administrators will be able to access and search all platform transactions from a centralized interface.",
        "6. Fraud Detection Analytics: The system will evaluate each processed transfer against configurable fraud risk factors, including transfer amounts above a defined threshold, multiple rapid transactions within a short time window, and transfers to newly added beneficiaries. Each transaction will be classified as LOW, MEDIUM, or HIGH risk, and a fraud log will be persisted and made available through customer and admin review interfaces.",
        "7. Financial Health Insights: The platform will generate a monthly financial health score for each customer by analyzing total credits, total debits, and savings rate. Monthly credit and debit trends along with personalized smart suggestions will be presented to help customers make better financial decisions.",
        "8. Administrative Dashboard: The admin module will provide centralized visibility into total registered customers, transaction count, high-risk transaction alerts, and aggregate transaction volume. Administrators will also be able to search customers and view all fraud logs across the platform.",
    ]),
    ("Development Methodology:", [
        "An Agile and iterative development approach will be followed throughout the project. Implementation will be organized into structured phases targeting authentication, core banking modules, transfer processing, fraud analytics, financial insights, and final integration. Git-based version control will be used for managing the source code. Functional testing, REST API validation through Swagger UI, security review of JWT flows, exception handling review, and frontend UI validation will be performed throughout the development lifecycle to improve system reliability, correctness, and usability.",
    ]),
    ("Limitation:", [
        "1. Mobile Application: The project will be limited to a responsive web application. Native mobile applications for Android or iOS will not be included in the current scope.",
        "2. Advanced Machine Learning Fraud Detection: The fraud detection system in this phase will be rule-based and AI-inspired, using configurable scoring logic rather than a fully trained machine learning model based on real banking transaction datasets.",
        "3. External Payment Integration: The platform will simulate internal transfers within the managed system. Integration with external inter-bank transfer systems such as NEFT, RTGS, or third-party banking gateways will not be included in this phase.",
        "4. Token Refresh and Logout Invalidation: The current implementation will use JWT-based authentication without refresh token support or server-side token invalidation on logout. These are identified as future enhancements.",
        "5. Scalability Infrastructure: Advanced production infrastructure features such as distributed deployment, load balancing, clustering, and cloud-native auto-scaling will remain outside the present project scope.",
        "6. Security Auditing: Formal penetration testing, PCI-DSS compliance certification, and enterprise-grade security audit integration will not be fully covered within the academic scope of this project.",
        "7. Automated Testing Coverage: Comprehensive unit tests and integration test suites for backend and frontend modules will not be fully included in the current phase and are identified as future improvements.",
    ]),
]

WORK_PLAN = [
    ("Week 1", "Conduct requirement analysis of digital banking workflows and finalize the project scope, user roles, and technology stack. Set up the Git repository, initialize the Spring Boot backend project with Maven, and scaffold the React frontend structure. Prepare initial technical documentation and system architecture diagrams."),
    ("Week 2", "Design the MySQL relational database schema for users, roles, accounts, beneficiaries, transactions, and fraud logs. Define REST API endpoints across authentication, customer, and admin modules. Implement JPA entity classes, repository interfaces, and configure Spring Boot datasource and Hibernate settings."),
    ("Week 3", "Implement Spring Security configuration with JWT authentication filter, custom user details service, and BCrypt password encoding. Develop user registration and login APIs with DTO validation, duplicate email and phone checks, and JWT token generation."),
    ("Week 4", "Develop customer-facing APIs for profile retrieval and update, account summary, and beneficiary CRUD operations. Implement transaction history APIs with search filtering by description and reference number. Build a global exception handler for consistent API error responses."),
    ("Week 5", "Implement the secure fund transfer module with atomic balance updates, dual-entry transaction recording, and fraud log generation. Develop the fraud detection service using risk scoring rules for large amounts, rapid consecutive transfers, and newly added beneficiaries. Build the financial health insight service with monthly trend analysis, health score calculation, and smart suggestions."),
    ("Week 6", "Develop frontend pages for login, registration, customer dashboard, transfer money, beneficiary management, and transaction history. Implement JWT-based authentication context, protected routing, Axios interceptors, and role-based navigation."),
    ("Week 7", "Develop and integrate fraud analytics and financial insights frontend pages with trend displays and suggestion cards. Build the admin dashboard with customer management, all-transactions view, and fraud log monitoring. Perform functional testing, API validation through Swagger UI, and cross-browser responsiveness review."),
    ("Week 8", "Conduct end-to-end system review covering authentication flows, transfer scenarios, fraud detection behavior, and role-based access restrictions. Perform debugging, exception handling review, security validation of JWT flows, and code cleanup. Finalize project report, synopsis, ER diagram documentation, and presentation materials for submission and viva."),
]


def set_times_new_roman(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold


def set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0):
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_times_new_roman(run, size=12)


def generate_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(12)

    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    r = p.add_run("Project Synopsis")
    set_times_new_roman(r, size=14, bold=True)

    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    r = p.add_run("Title")
    set_times_new_roman(r, size=14, bold=True)

    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, after=12)
    r = p.add_run(TITLE)
    set_times_new_roman(r, size=12)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in DETAILS:
        row = table.add_row()
        for idx, text in enumerate((label, value)):
            cell_para = row.cells[idx].paragraphs[0]
            set_paragraph_format(cell_para, WD_ALIGN_PARAGRAPH.LEFT)
            run = cell_para.add_run(text)
            set_times_new_roman(run, size=12)

    doc.add_paragraph("")

    for heading, paragraphs in SECTIONS:
        p = doc.add_paragraph()
        set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
        r = p.add_run(heading)
        set_times_new_roman(r, size=12, bold=True)
        for text in paragraphs:
            p = doc.add_paragraph()
            set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=3)
            r = p.add_run(text)
            set_times_new_roman(r, size=12)

    doc.add_page_break()

    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, after=8)
    r = p.add_run("Work Plan (Week 1 to Week 8)")
    set_times_new_roman(r, size=14, bold=True)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = ["Week No.", "Activities to be Completed"]
    for idx, text in enumerate(headers):
        para = table.rows[0].cells[idx].paragraphs[0]
        set_paragraph_format(para, WD_ALIGN_PARAGRAPH.LEFT)
        run = para.add_run(text)
        set_times_new_roman(run, size=12, bold=True)

    for week, activity in WORK_PLAN:
        row = table.add_row()
        for idx, text in enumerate((week, activity)):
            para = row.cells[idx].paragraphs[0]
            set_paragraph_format(para, WD_ALIGN_PARAGRAPH.LEFT)
            run = para.add_run(text)
            set_times_new_roman(run, size=12)

    doc.save(DOCX_PATH)


def build_pdf():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Title14", fontName="Times-Bold", fontSize=14, leading=18, alignment=TA_CENTER, spaceAfter=8))
    styles.add(ParagraphStyle(name="Body12", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_JUSTIFY, spaceAfter=6))
    styles.add(ParagraphStyle(name="Center12", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_CENTER, spaceAfter=6))
    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm, topMargin=25 * mm, bottomMargin=25 * mm)
    story = []

    story.append(Paragraph("Project Synopsis", styles["Title14"]))
    story.append(Paragraph("Title", styles["Title14"]))
    story.append(Paragraph(TITLE, styles["Center12"]))
    story.append(Spacer(1, 8))

    details_table = Table([[a, b] for a, b in DETAILS], colWidths=[55 * mm, 95 * mm])
    details_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.8, colors.black),
        ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("LEADING", (0, 0), (-1, -1), 18),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(details_table)
    story.append(Spacer(1, 10))

    for heading, paragraphs in SECTIONS:
        story.append(Paragraph(f"<b>{heading}</b>", styles["Body12"]))
        for text in paragraphs:
            story.append(Paragraph(text, styles["Body12"]))

    story.append(PageBreak())
    story.append(Paragraph("Work Plan (Week 1 to Week 8)", styles["Title14"]))
    data = [["Week No.", "Activities to be Completed"]] + [[w, a] for w, a in WORK_PLAN]
    work_table = Table(data, colWidths=[28 * mm, 122 * mm])
    work_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.8, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("LEADING", (0, 0), (-1, -1), 18),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(work_table)

    def page_number(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Times-Roman", 12)
        page_width, _ = A4
        canvas.drawCentredString(page_width / 2, 12 * mm, str(doc_obj.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)


if __name__ == "__main__":
    generate_docx()
    build_pdf()
