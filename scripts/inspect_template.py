import docx
import glob
import os

files = glob.glob("*Template*.docx")
if not files:
    print("No template found.")
    exit(1)

doc = docx.Document(files[0])
os.makedirs("C:/Users/bvche/.gemini/antigravity/brain/fddfed3d-5023-4a51-a682-451aa3ebd913/scratch", exist_ok=True)
with open("C:/Users/bvche/.gemini/antigravity/brain/fddfed3d-5023-4a51-a682-451aa3ebd913/scratch/template_structure.txt", "w", encoding="utf-8") as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"P {i}: style={p.style.name}, text='{p.text}'\n")
        
    f.write("\n=== TABLES ===\n")
    for i, t in enumerate(doc.tables):
        f.write(f"T {i}: {len(t.rows)}x{len(t.columns)}\n")
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                f.write(f"  Cell ({r_idx},{c_idx}): '{cell.text.replace(chr(10), chr(32))}'\n")

print("Done. Saved to template_structure.txt")
